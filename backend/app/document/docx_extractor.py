from __future__ import annotations

from pathlib import Path

from docx import Document


def extract_docx_text(file_path: str | Path) -> tuple[str, float]:
    """
    Extract text from DOCX paragraphs and tables.

    Returns:
        (extracted_text, confidence)
    """
    path = Path(file_path)

    document = Document(str(path))

    sections: list[str] = []

    # Paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            sections.append(text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            cells = [cell for cell in cells if cell]

            if cells:
                sections.append(" | ".join(cells))

    extracted_text = "\n".join(sections).strip()

    if not extracted_text:
        return "", 0.0

    confidence = 0.95

    return extracted_text, confidence